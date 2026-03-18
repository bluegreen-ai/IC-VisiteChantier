#!/usr/bin/env python3
"""
Render CR Visite from template + context JSON + photos.

Usage:
    python render_cr_visite.py context.json [--photos-dir ./photos] [--output cr.docx]

Context JSON format:
{
    "titre_service": "Suivi de réfection des balcons - Lot 12",
    "client": "SDC Le Gros Saule",
    "residence": "Résidence Savigny Impair",
    "batiments_visites": "Bâtiment A",
    "adresse": "1-50 avenue de Savigny",
    "code_postal_ville": "93600 Aulnay-sous-Bois",
    "ref_dossier": "DE0328",
    "date_visite": "27 février 2026",
    "participants": [
        {"nom": "R. Laborbe", "fonction": "M.O Lot 12", "entreprise": "IC Ingénieurs Conseils", "contact": "06 50 96 61 98"},
        {"nom": "A. Plassart", "fonction": "Ingénieur Travaux", "entreprise": "Bouygues Bâtiment", "contact": "06 66 87 96 80"}
    ],
    "objet_visite": "IC Ingénieurs Conseils assure le suivi...",
    "synthese": "L'inspection a permis...",
    "observations": [
        {"ref": "V1-01", "etage_facade": "10ème — Façade Est", "observation": "Traces de truelle", "action": "Reprendre les traces", "photo": "photo1.jpg"},
        ...
    ],
    "conclusion": "La visite confirme..."
}
"""

import json
import os
import sys
import argparse
from copy import deepcopy
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from PIL import Image

MOIS_FR = [
    'janvier', 'février', 'mars', 'avril', 'mai', 'juin',
    'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre',
]


def format_date_french(iso_date: str) -> str:
    """Convert ISO date (YYYY-MM-DD) to French format (D mois YYYY).

    Passes through non-ISO strings unchanged for backward compatibility.
    """
    try:
        dt = datetime.strptime(iso_date, "%Y-%m-%d")
        return f"{dt.day} {MOIS_FR[dt.month - 1]} {dt.year}"
    except (ValueError, TypeError):
        return iso_date


# === Template path ===
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(SCRIPT_DIR, "template_cr_visite_aulnay.docx")


def set_cell_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:start w:w="{left}" w:type="dxa"/>'
        f'  <w:end w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_photo_to_cell(cell, photo_path, max_width_cm=5.0):
    if not os.path.exists(photo_path):
        cell.paragraphs[0].add_run("[Photo non disponible]").font.size = Pt(7)
        return False
    with Image.open(photo_path) as img:
        w, h = img.size
    aspect = h / w
    width = Cm(max_width_cm)
    max_h = Cm(6.5)
    if Cm(max_width_cm * aspect) > max_h:
        width = Cm(max_h.cm / aspect)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(photo_path, width=width)
    return True


def set_table_borders(table, color="AAAAAA"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl.insert(0, tblPr)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def render_cr(context, photos_dir=".", output_path="cr_visite.docx", template_path=None):
    """
    Render a CR Visite document from context data.

    Args:
        context: dict with all template variables
        photos_dir: directory containing observation photos
        output_path: path for the generated .docx
        template_path: path to the template (defaults to bundled template)
    """
    if template_path is None:
        template_path = TEMPLATE_PATH

    doc = Document(template_path)

    # === 1. Replace cover page variables ===
    var_map = {
        "{{ titre_service }}": context.get("titre_service", ""),
        "{{ client }}": context.get("client", ""),
        "{{ residence }}": context.get("residence", ""),
        "{{ batiments_visites }}": context.get("batiments_visites", ""),
        "{{ adresse }}": context.get("adresse", ""),
        "{{ code_postal_ville }}": context.get("code_postal_ville", ""),
        "{{ ref_dossier }}": context.get("ref_dossier", ""),
        "{{ date_visite }}": format_date_french(context.get("date_visite", "")),
        "{{ objet_visite }}": context.get("objet_visite", ""),
        "{{ synthese }}": context.get("synthese", ""),
        "{{ conclusion }}": context.get("conclusion", ""),
    }

    for p in doc.paragraphs:
        for key, value in var_map.items():
            if key in p.text:
                for run in p.runs:
                    if key in run.text:
                        run.text = run.text.replace(key, value)

    # Also replace date in intro paragraph of observations section
    date_french = format_date_french(context.get("date_visite", ""))
    for p in doc.paragraphs:
        if "observations relevées lors de la visite" in p.text:
            for run in p.runs:
                if "27 février 2026" in run.text:
                    run.text = run.text.replace("27 février 2026", date_french)
                elif "{{ date_visite }}" in run.text:
                    run.text = run.text.replace("{{ date_visite }}", date_french)

    # === 2. Populate participants table ===
    participants = context.get("participants", [])
    if participants and len(doc.tables) > 0:
        ptable = doc.tables[0]

        # Clear existing data rows
        while len(ptable.rows) > 1:
            tr = ptable.rows[-1]._tr
            ptable._tbl.remove(tr)

        # Add participant rows
        for p_data in participants:
            row = ptable.add_row()
            values = [
                p_data.get("nom", ""),
                p_data.get("fonction", ""),
                p_data.get("entreprise", ""),
                p_data.get("contact", "")
            ]
            for i, (cell, val) in enumerate(zip(row.cells, values)):
                cell.text = ""
                run = cell.paragraphs[0].add_run(val)
                run.font.size = Pt(9)
                set_cell_margins(cell, 30, 30, 60, 60)

    # === 3. Populate observations table ===
    observations = context.get("observations", [])
    if observations and len(doc.tables) > 1:
        obs_table = doc.tables[1]

        # Clear existing data rows (keep header)
        while len(obs_table.rows) > 1:
            tr = obs_table.rows[-1]._tr
            obs_table._tbl.remove(tr)

        # Add observation rows
        for idx, obs in enumerate(observations):
            row = obs_table.add_row()
            cells = row.cells

            # Col 1: Ref + Étage/Façade
            cells[0].text = ""
            p0 = cells[0].paragraphs[0]
            run_ref = p0.add_run(f"{obs.get('ref', '')}\n")
            run_ref.bold = True
            run_ref.font.size = Pt(8)
            run_ref.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            run_loc = p0.add_run(obs.get("etage_facade", ""))
            run_loc.font.size = Pt(9)
            set_cell_margins(cells[0], 50, 50, 80, 80)

            # Col 2: Observation + Action
            cells[1].text = ""
            p1 = cells[1].paragraphs[0]
            observation_text = obs.get("observation", "")
            action_text = obs.get("action", "")

            run_obs = p1.add_run(observation_text)
            run_obs.font.size = Pt(9)

            if action_text:
                p1.add_run("\n")
                run_act = p1.add_run(f"→ {action_text}")
                run_act.font.size = Pt(8)
                run_act.italic = True
                run_act.font.color.rgb = RGBColor(0x80, 0x40, 0x00)

            set_cell_margins(cells[1], 50, 50, 80, 80)

            # Col 3: Photo
            cells[2].text = ""
            photo_file = obs.get("photo", "")
            if photo_file:
                photo_path = os.path.join(photos_dir, photo_file)
                add_photo_to_cell(cells[2], photo_path, max_width_cm=5.0)
            set_cell_margins(cells[2], 30, 30, 40, 40)

            # Alternate row shading
            if idx % 2 == 0:
                for cell in cells:
                    set_cell_shading(cell, "F0F4F8")

        set_table_borders(obs_table)

    # === 4. Save ===
    doc.save(output_path)
    print(f"✅ CR Visite saved: {output_path}")
    print(f"   {len(participants)} participants, {len(observations)} observations")
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Render CR Visite from template")
    parser.add_argument("context", help="Path to context JSON file")
    parser.add_argument("--photos-dir", default=".", help="Directory containing photos")
    parser.add_argument("--output", default="cr_visite.docx", help="Output path")
    parser.add_argument("--template", default=None, help="Template path (optional)")
    args = parser.parse_args()

    with open(args.context, 'r', encoding='utf-8') as f:
        ctx = json.load(f)

    render_cr(ctx, args.photos_dir, args.output, args.template)
